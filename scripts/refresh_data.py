#!/usr/bin/env python3
"""Refresh the National Venturing Dashboard from the newest local workbook capture."""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "source"
DATA_DIR = ROOT / "data"
JSON_OUTPUT = DATA_DIR / "dashboard-data.json"
CSV_OUTPUT = DATA_DIR / "crews.csv"
TZ = ZoneInfo("America/Chicago")

CSV_FIELDS = (
    "unitId",
    "cst",
    "council",
    "district",
    "crewNumber",
    "crewName",
    "renewalDate",
    "renewalWorkflow",
    "youth",
    "adults",
    "primaryYouth",
    "primaryYoY",
    "primaryToYearEnd",
    "unitMetric",
    "sizeHealthy",
    "growthHealthy",
    "leadershipHealthy",
    "outdoorHealthy",
    "unitLeaderTrained",
    "committeeChairTrained",
    "officersTrained",
    "retentionHealthy",
)


def clean(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return " ".join(value.strip().split())
    return value


def iso_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = clean(value)
    return str(text)[:10] if text else ""


def yes(value):
    if value is None:
        return None
    return str(clean(value)).casefold() == "yes"


def number(value, default=0):
    return value if isinstance(value, (int, float)) else default


def rows_by_header(workbook, sheet_name):
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(iterator)]
    lookup = {header: index for index, header in enumerate(headers) if header}
    return lookup, iterator


def newest_source():
    candidates = [path for path in SOURCE_DIR.glob("Venturing*.xlsx") if not path.name.startswith("~$")]
    if not candidates:
        raise FileNotFoundError(f"No Venturing workbook was found in {SOURCE_DIR}")
    return max(candidates, key=lambda path: (path.stat().st_mtime, path.name))


def workbook_as_of(workbook):
    value = workbook["Dashboard"]["B2"].value
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    raise ValueError("Dashboard!B2 does not contain the expected workbook snapshot date")


def build(source):
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    required_sheets = {"Dashboard", "Crews", "UnitMetric"}
    missing_sheets = required_sheets.difference(workbook.sheetnames)
    if missing_sheets:
        raise ValueError(f"Workbook is missing required sheets: {', '.join(sorted(missing_sheets))}")

    crew_header, crew_rows = rows_by_header(workbook, "Crews")
    metric_header, metric_rows = rows_by_header(workbook, "UnitMetric")
    required_crew = {
        "Unit_ID", "CST", "Council", "District", "Crew", "Renewal", "Total Youth",
        "Total Adults", "Crew Name", "Primary Youth", "Primary YOY", "Primary to YE",
        "Unit Metric", "Renewal Status",
    }
    required_metric = {
        "UnitID", "UnitMetric", "Size", "Growth", "Youth Leadership", "Outdoor",
        "UL Trained", "CC Trained", "Officers Trained", "Retention",
    }
    if required_crew.difference(crew_header):
        raise ValueError(f"Crews sheet is missing columns: {', '.join(sorted(required_crew.difference(crew_header)))}")
    if required_metric.difference(metric_header):
        raise ValueError(f"UnitMetric sheet is missing columns: {', '.join(sorted(required_metric.difference(metric_header)))}")

    metric_by_unit = {
        str(clean(row[metric_header["UnitID"]])): row
        for row in metric_rows
        if clean(row[metric_header["UnitID"]]) != ""
    }

    crews = []
    for row in crew_rows:
        unit_id = clean(row[crew_header["Unit_ID"]])
        if unit_id == "":
            continue
        metric = metric_by_unit.get(str(unit_id))

        def metric_value(name, default=None):
            if metric is None:
                return default
            return clean(metric[metric_header[name]])

        unit_metric = number(row[crew_header["Unit Metric"]], number(metric_value("UnitMetric"), None))
        crews.append(
            {
                "unitId": str(unit_id),
                "cst": str(clean(row[crew_header["CST"]])),
                "council": clean(row[crew_header["Council"]]),
                "district": clean(row[crew_header["District"]]),
                "crewNumber": str(clean(row[crew_header["Crew"]])),
                "crewName": clean(row[crew_header["Crew Name"]]),
                "renewalDate": iso_date(row[crew_header["Renewal"]]),
                "renewalWorkflow": clean(row[crew_header["Renewal Status"]]) or "Not in renewal workflow",
                "youth": number(row[crew_header["Total Youth"]], None),
                "adults": number(row[crew_header["Total Adults"]], None),
                "primaryYouth": number(row[crew_header["Primary Youth"]], None),
                "primaryYoY": number(row[crew_header["Primary YOY"]], None),
                "primaryToYearEnd": number(row[crew_header["Primary to YE"]], None),
                "unitMetric": unit_metric,
                "sizeHealthy": yes(metric_value("Size")),
                "growthHealthy": yes(metric_value("Growth")),
                "leadershipHealthy": yes(metric_value("Youth Leadership")),
                "outdoorHealthy": yes(metric_value("Outdoor")),
                "unitLeaderTrained": yes(metric_value("UL Trained")),
                "committeeChairTrained": yes(metric_value("CC Trained")),
                "officersTrained": yes(metric_value("Officers Trained")),
                "retentionHealthy": yes(metric_value("Retention")),
            }
        )

    unit_ids = [crew["unitId"] for crew in crews]
    duplicate_ids = [unit_id for unit_id, count in Counter(unit_ids).items() if count > 1]
    if duplicate_ids:
        raise ValueError(f"Crews sheet contains duplicate Unit IDs: {', '.join(duplicate_ids[:10])}")

    matched = sum(1 for crew in crews if crew["unitId"] in metric_by_unit)
    snapshot = {
        "meta": {
            "title": "National Venturing Dashboard",
            "asOf": workbook_as_of(workbook),
            "generatedAt": datetime.now(TZ).isoformat(timespec="seconds"),
            "source": "Newest local Venturing workbook capture",
            "sourceCaptureName": source.name,
            "rowCount": len(crews),
            "councilCount": len({crew["council"] for crew in crews}),
            "cstCount": len({crew["cst"] for crew in crews}),
            "metricJoinMatched": matched,
            "metricJoinMissing": len(crews) - matched,
            "missingPrimaryYouth": sum(crew["primaryYouth"] is None for crew in crews),
            "missingUnitMetric": sum(crew["unitMetric"] is None for crew in crews),
            "privacy": "Contact names, email addresses, phone numbers, meeting addresses, websites, and charter-organization fields are excluded from the dashboard data.",
        },
        "crews": crews,
    }

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    JSON_OUTPUT.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    with CSV_OUTPUT.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=CSV_FIELDS)
        writer.writeheader()
        writer.writerows(crews)
    return snapshot


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, help="Workbook to use instead of the newest source/Venturing*.xlsx capture")
    args = parser.parse_args()
    source = (args.source or newest_source()).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    snapshot = build(source)
    result = {
        "source": source.name,
        "asOf": snapshot["meta"]["asOf"],
        "crews": snapshot["meta"]["rowCount"],
        "councils": snapshot["meta"]["councilCount"],
        "csts": snapshot["meta"]["cstCount"],
        "metricJoinMatched": snapshot["meta"]["metricJoinMatched"],
        "metricJoinMissing": snapshot["meta"]["metricJoinMissing"],
        "generatedFiles": [str(CSV_OUTPUT.relative_to(ROOT)), str(JSON_OUTPUT.relative_to(ROOT))],
    }
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
